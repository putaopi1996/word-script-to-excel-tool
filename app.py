from __future__ import annotations

import json
import time
from io import BytesIO
from pathlib import Path

import requests as http_client
from flask import Flask, jsonify, render_template, request, send_file

from services.exporter import export_rows_to_excel, export_table_to_excel
from services.parser import ScriptRow, parse_document


CONFIG_DIR = Path.home() / ".word-script-excel-tool"
CONFIG_FILE = CONFIG_DIR / "config.json"


def _get_config_path() -> Path:
    """Get config file path next to the running exe/script."""
    import sys
    if getattr(sys, "frozen", False):
        # Running as packaged exe
        base = Path(sys.executable).resolve().parent
    else:
        # Running as script
        base = Path(__file__).resolve().parent
    return base / "config.json"


def load_config() -> dict:
    config_file = _get_config_path()
    if config_file.exists():
        try:
            return json.loads(config_file.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            return {}
    return {}


def save_config(data: dict) -> None:
    config_file = _get_config_path()
    config_file.parent.mkdir(parents=True, exist_ok=True)
    config_file.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def build_output_filename(filename: str) -> str:
    source_name = Path(filename).stem or "\u5267\u672c"
    return f"{source_name}-\u914d\u8868.xlsx"


def serialize_rows(rows: list[ScriptRow]) -> list[dict[str, str]]:
    return [
        {
            "speaker": row.speaker,
            "content": row.content,
            "original_text": row.original_text,
        }
        for row in rows
    ]


def deserialize_rows(payload: list[dict[str, object]] | None) -> list[ScriptRow]:
    if not payload:
        raise ValueError("\u6ca1\u6709\u53ef\u5bfc\u51fa\u7684\u9884\u89c8\u6570\u636e")

    rows: list[ScriptRow] = []
    for item in payload:
        rows.append(
            ScriptRow(
                speaker=str(item.get("speaker", "")).strip(),
                content=str(item.get("content", "")).strip(),
                original_text=str(item.get("original_text", "")).strip(),
            )
        )

    return rows


def deserialize_table_rows(payload: list[dict[str, object]] | None) -> list[dict[str, str]]:
    if not payload:
        raise ValueError("\u6ca1\u6709\u53ef\u5bfc\u51fa\u7684\u9884\u89c8\u6570\u636e")

    rows: list[dict[str, str]] = []
    for item in payload:
        rows.append({str(key): str(value or "").strip() for key, value in item.items()})
    return rows


def deserialize_columns(payload: list[dict[str, object]] | None) -> list[dict[str, str]]:
    if not payload:
        raise ValueError("\u6ca1\u6709\u53ef\u5bfc\u51fa\u7684\u5217\u914d\u7f6e")

    columns: list[dict[str, str]] = []
    for item in payload:
        key = str(item.get("key", "")).strip()
        label = str(item.get("label", "")).strip()
        if not key or not label:
            continue
        columns.append({"key": key, "label": label})

    if not columns:
        raise ValueError("\u6ca1\u6709\u53ef\u5bfc\u51fa\u7684\u5217\u914d\u7f6e")

    return columns


def create_app(resource_root: Path | None = None) -> Flask:
    root = Path(resource_root) if resource_root else Path(__file__).resolve().parent
    app = Flask(
        __name__,
        template_folder=str(root / "templates"),
        static_folder=str(root / "static"),
    )
    app.config["MAX_CONTENT_LENGTH"] = 100 * 1024 * 1024  # 100MB

    @app.get("/")
    def index():
        return render_template("index.html")

    @app.get("/api/version")
    def get_version():
        changelog_file = root / "changelog.json"
        try:
            data = json.loads(changelog_file.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            data = {"version": "unknown", "entries": []}
        return jsonify(data)

    @app.post("/api/convert")
    def convert():
        upload = request.files.get("file")
        if upload is None or not upload.filename:
            return jsonify({"error": "\u8bf7\u5148\u9009\u62e9\u4e00\u4e2a Word \u6587\u4ef6"}), 400

        if Path(upload.filename).suffix.lower() != ".docx":
            return jsonify({"error": "\u4ec5\u652f\u6301\u4e0a\u4f20 .docx \u6587\u4ef6"}), 400

        include_original = request.form.get("include_original", "false").lower() == "true"

        try:
            rows = parse_document(upload.stream)
            content = export_rows_to_excel(rows, include_original=include_original)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400

        return send_file(
            BytesIO(content),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=build_output_filename(upload.filename),
        )

    @app.post("/api/preview")
    def preview():
        upload = request.files.get("file")
        if upload is None or not upload.filename:
            return jsonify({"error": "\u8bf7\u5148\u9009\u62e9\u4e00\u4e2a Word \u6587\u4ef6"}), 400

        if Path(upload.filename).suffix.lower() != ".docx":
            return jsonify({"error": "\u4ec5\u652f\u6301\u4e0a\u4f20 .docx \u6587\u4ef6"}), 400

        try:
            rows = parse_document(upload.stream)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400

        return jsonify({"rows": serialize_rows(rows)})

    @app.post("/api/export-rows")
    def export_rows():
        payload = request.get_json(silent=True) or {}
        source_filename = str(payload.get("source_filename", "\u5267\u672c.docx"))

        try:
            columns = deserialize_columns(payload.get("columns"))
            rows = deserialize_table_rows(payload.get("rows"))
            content = export_table_to_excel(columns, rows)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400

        return send_file(
            BytesIO(content),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=build_output_filename(source_filename),
        )

    @app.post("/api/ai/models")
    def ai_models():
        payload = request.get_json(silent=True) or {}
        api_url = str(payload.get("api_url", "")).strip().rstrip("/")
        api_key = str(payload.get("api_key", "")).strip()

        if not api_url or not api_key:
            return jsonify({"error": "请填写 API 地址和密钥"}), 400

        try:
            resp = http_client.get(
                f"{api_url}/v1/models",
                headers={"Authorization": f"Bearer {api_key}"},
                timeout=15,
            )
            resp.raise_for_status()
            data = resp.json()
            models = [m.get("id", "") for m in data.get("data", []) if m.get("id")]
            return jsonify({"models": sorted(models)})
        except http_client.Timeout:
            return jsonify({"error": "请求超时，请检查 API 地址"}), 502
        except http_client.RequestException as exc:
            return jsonify({"error": f"请求失败: {exc}"}), 502
        except (ValueError, KeyError):
            return jsonify({"error": "返回数据格式异常"}), 502

    @app.post("/api/ai/test")
    def ai_test():
        payload = request.get_json(silent=True) or {}
        api_url = str(payload.get("api_url", "")).strip().rstrip("/")
        api_key = str(payload.get("api_key", "")).strip()
        model = str(payload.get("model", "")).strip()

        if not api_url or not api_key:
            return jsonify({"error": "请填写 API 地址和密钥"}), 400

        if not model:
            return jsonify({"error": "请填写模型名称"}), 400

        try:
            start = time.time()
            resp = http_client.post(
                f"{api_url}/v1/chat/completions",
                headers={
                    "Authorization": f"Bearer {api_key}",
                    "Content-Type": "application/json",
                },
                json={
                    "model": model,
                    "messages": [{"role": "user", "content": "Hi"}],
                    "max_tokens": 5,
                },
                timeout=30,
            )
            elapsed = round((time.time() - start) * 1000)
            resp.raise_for_status()
            return jsonify({"success": True, "latency_ms": elapsed})
        except http_client.Timeout:
            return jsonify({"error": "请求超时"}), 502
        except http_client.RequestException as exc:
            return jsonify({"error": f"连接失败: {exc}"}), 502

    @app.get("/api/config")
    def get_config():
        return jsonify(load_config())

    @app.post("/api/config")
    def post_config():
        payload = request.get_json(silent=True) or {}
        config = load_config()
        config.update(payload)
        save_config(config)
        return jsonify({"success": True})

    @app.post("/api/ai/check-typos")
    def ai_check_typos():
        payload = request.get_json(silent=True) or {}
        rows = payload.get("rows", [])
        prompt_template = str(payload.get("prompt", "")).strip()

        if not rows:
            return jsonify({"error": "没有可检测的内容"}), 400

        config = load_config()
        api_url = str(config.get("ai_api_url", "")).strip().rstrip("/")
        api_key = str(config.get("ai_api_key", "")).strip()
        model = str(config.get("ai_model", "")).strip()

        if not api_url or not api_key or not model:
            return jsonify({"error": "请先在设置中配置 AI API"}), 400

        # Build content lines
        lines = []
        for i, row in enumerate(rows, 1):
            speaker = str(row.get("speaker", "")).strip()
            content = str(row.get("content", "")).strip()
            lines.append(f"[{i}] 角色：{speaker} | 台词：{content}")

        content_text = "\n".join(lines)

        # Build final prompt
        if prompt_template and "{{content}}" in prompt_template:
            final_prompt = prompt_template.replace("{{content}}", content_text)
        else:
            final_prompt = prompt_template + "\n\n" + content_text if prompt_template else content_text

        try:
            resp = http_client.post(
                f"{api_url}/v1/chat/completions",
                headers={
                    "Authorization": f"Bearer {api_key}",
                    "Content-Type": "application/json",
                },
                json={
                    "model": model,
                    "messages": [{"role": "user", "content": final_prompt}],
                    "temperature": 0.1,
                },
                timeout=120,
            )
            resp.raise_for_status()
            data = resp.json()
        except http_client.Timeout:
            return jsonify({"error": "AI 请求超时，内容可能过长"}), 502
        except http_client.RequestException as exc:
            return jsonify({"error": f"AI 请求失败: {exc}"}), 502

        # Parse AI response
        try:
            ai_content = data["choices"][0]["message"]["content"].strip()
            # Try to extract JSON from response (handle markdown code blocks)
            if ai_content.startswith("```"):
                # Remove markdown code block wrapper
                ai_content = ai_content.split("\n", 1)[1] if "\n" in ai_content else ai_content[3:]
                if ai_content.endswith("```"):
                    ai_content = ai_content[:-3].strip()
            issues = json.loads(ai_content)
            if not isinstance(issues, list):
                issues = []
        except (KeyError, IndexError, json.JSONDecodeError):
            return jsonify({"error": "AI 返回格式异常，请重试"}), 502

        return jsonify({"issues": issues})

    # ===== SFX Dictionary & Matching APIs =====

    @app.get("/api/sfx-dict")
    def get_sfx_dict():
        config = load_config()
        return jsonify({"dictionary": config.get("sfx_dictionary", [])})

    @app.post("/api/sfx-dict")
    def post_sfx_dict():
        payload = request.get_json(silent=True) or {}
        dictionary = payload.get("dictionary", [])

        # Validate: list of {name, id}
        cleaned: list[dict[str, str]] = []
        for item in dictionary:
            name = str(item.get("name", "")).strip()
            sfx_id = str(item.get("id", "")).strip()
            if name or sfx_id:
                cleaned.append({"name": name, "id": sfx_id})

        # Check for duplicate names
        seen: dict[str, list[int]] = {}
        for i, entry in enumerate(cleaned):
            key = entry["name"]
            if key in seen:
                seen[key].append(i)
            else:
                seen[key] = [i]

        duplicates = {k: v for k, v in seen.items() if len(v) > 1}
        if duplicates:
            return jsonify({
                "error": "存在重复的音效名称，无法保存",
                "duplicates": duplicates,
            }), 400

        config = load_config()
        config["sfx_dictionary"] = cleaned
        save_config(config)
        return jsonify({"success": True, "count": len(cleaned)})

    @app.post("/api/sfx-match/import")
    def sfx_match_import():
        upload = request.files.get("file")
        if upload is None or not upload.filename:
            return jsonify({"error": "请选择一个 Excel 文件"}), 400

        if Path(upload.filename).suffix.lower() not in (".xlsx", ".xls"):
            return jsonify({"error": "仅支持上传 .xlsx 文件"}), 400

        try:
            from openpyxl import load_workbook

            data = upload.stream.read()
            wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
            ws = wb.active

            # Read headers (first row)
            headers: list[str] = []
            rows_data: list[list[str]] = []

            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    headers = [str(cell or "").strip() for cell in row]
                else:
                    rows_data.append([str(cell or "").strip() for cell in row])

            wb.close()

            if not headers:
                return jsonify({"error": "Excel 文件为空"}), 400

            return jsonify({
                "headers": headers,
                "rows": rows_data,
                "filename": upload.filename,
            })
        except Exception as exc:
            return jsonify({"error": f"解析 Excel 失败: {exc}"}), 400

    @app.post("/api/sfx-match/export")
    def sfx_match_export():
        payload = request.get_json(silent=True) or {}
        rows = payload.get("rows", [])
        source_filename = str(payload.get("source_filename", "音效匹配.xlsx"))

        if not rows:
            return jsonify({"error": "没有可导出的数据"}), 400

        columns = [
            {"key": "sfx_id", "label": "音效ID"},
            {"key": "sfx_name", "label": "文件名称"},
        ]

        try:
            content = export_table_to_excel(columns, rows)
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400

        out_name = Path(source_filename).stem + "-已匹配.xlsx"
        return send_file(
            BytesIO(content),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=out_name,
        )

    @app.post("/api/ai/generate-sfx")
    def ai_generate_sfx():
        payload = request.get_json(silent=True) or {}
        rows = payload.get("rows", [])
        prompt_template = str(payload.get("prompt", "")).strip()

        if not rows:
            return jsonify({"error": "没有可生成的内容"}), 400

        config = load_config()
        api_url = str(config.get("ai_api_url", "")).strip().rstrip("/")
        api_key = str(config.get("ai_api_key", "")).strip()
        model = str(config.get("ai_model", "")).strip()

        if not api_url or not api_key or not model:
            return jsonify({"error": "请先在设置中配置 AI API"}), 400

        lines = []
        for i, row in enumerate(rows, 1):
            speaker = str(row.get("speaker", "")).strip()
            content = str(row.get("content", "")).strip()
            lines.append(f"[{i}] 角色：{speaker} | 台词：{content}")

        content_text = "\n".join(lines)

        if prompt_template and "{{content}}" in prompt_template:
            final_prompt = prompt_template.replace("{{content}}", content_text)
        else:
            final_prompt = prompt_template + "\n\n" + content_text if prompt_template else content_text

        try:
            resp = http_client.post(
                f"{api_url}/v1/chat/completions",
                headers={
                    "Authorization": f"Bearer {api_key}",
                    "Content-Type": "application/json",
                },
                json={
                    "model": model,
                    "messages": [{"role": "user", "content": final_prompt}],
                    "temperature": 0.3,
                },
                timeout=120,
            )
            resp.raise_for_status()
            data = resp.json()
        except http_client.Timeout:
            return jsonify({"error": "AI 请求超时，内容可能过长"}), 502
        except http_client.RequestException as exc:
            return jsonify({"error": f"AI 请求失败: {exc}"}), 502

        try:
            ai_content = data["choices"][0]["message"]["content"].strip()
            if ai_content.startswith("```"):
                ai_content = ai_content.split("\n", 1)[1] if "\n" in ai_content else ai_content[3:]
                if ai_content.endswith("```"):
                    ai_content = ai_content[:-3].strip()
            sfx_list = json.loads(ai_content)
            if not isinstance(sfx_list, list):
                sfx_list = []
        except (KeyError, IndexError, json.JSONDecodeError):
            return jsonify({"error": "AI 返回格式异常，请重试"}), 502

        return jsonify({"sfx": sfx_list})

    return app


app = create_app()


if __name__ == "__main__":
    app.run(debug=True)
