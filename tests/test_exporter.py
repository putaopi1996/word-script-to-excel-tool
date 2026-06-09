from io import BytesIO

from openpyxl import load_workbook

from services.exporter import export_rows_to_excel
from services.parser import ScriptRow


def sample_rows() -> list[ScriptRow]:
    return [
        ScriptRow(
            original_text="\u3010\u767d\u9053\u957f\u3011\u4f60\u6765\u4e86",
            speaker="\u767d\u9053\u957f",
            content="\u4f60\u6765\u4e86",
        ),
        ScriptRow(
            original_text="\u573a\u666f\u4e00 \u591c \u5916",
            speaker="\u65c1\u767d",
            content="\u573a\u666f\u4e00 \u591c \u5916",
        ),
    ]


def read_sheet_values(content: bytes) -> list[tuple]:
    workbook = load_workbook(filename=BytesIO(content))
    sheet = workbook.active
    return list(sheet.iter_rows(values_only=True))


def test_export_rows_to_excel_without_original_column() -> None:
    content = export_rows_to_excel(sample_rows(), include_original=False)

    values = read_sheet_values(content)

    assert values[0] == ("\u89d2\u8272", "\u53f0\u8bcd")
    assert values[1] == ("\u767d\u9053\u957f", "\u4f60\u6765\u4e86")
    assert values[2] == ("\u65c1\u767d", "\u573a\u666f\u4e00 \u591c \u5916")


def test_export_rows_to_excel_with_original_column() -> None:
    content = export_rows_to_excel(sample_rows(), include_original=True)

    values = read_sheet_values(content)

    assert values[0] == ("\u89d2\u8272", "\u53f0\u8bcd", "\u539f\u6587")
    assert values[1] == ("\u767d\u9053\u957f", "\u4f60\u6765\u4e86", "\u3010\u767d\u9053\u957f\u3011\u4f60\u6765\u4e86")
    assert values[2] == ("\u65c1\u767d", "\u573a\u666f\u4e00 \u591c \u5916", "\u573a\u666f\u4e00 \u591c \u5916")
