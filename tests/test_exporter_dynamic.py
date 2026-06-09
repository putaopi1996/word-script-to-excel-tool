from io import BytesIO

from openpyxl import load_workbook

from services.exporter import export_table_to_excel


def read_sheet_values(content: bytes) -> list[tuple]:
    workbook = load_workbook(filename=BytesIO(content))
    sheet = workbook.active
    return list(sheet.iter_rows(values_only=True))


def test_export_table_to_excel_uses_dynamic_columns_and_values() -> None:
    content = export_table_to_excel(
        columns=[
            {"key": "speaker", "label": "\u89d2\u8272"},
            {"key": "mood", "label": "\u60c5\u7eea"},
            {"key": "content", "label": "\u53f0\u8bcd"},
        ],
        rows=[
            {
                "speaker": "\u767d\u9053\u957f",
                "mood": "\u51b7\u9759",
                "content": "\u4f60\u6765\u4e86",
            }
        ],
    )

    values = read_sheet_values(content)

    assert values[0] == ("\u89d2\u8272", "\u60c5\u7eea", "\u53f0\u8bcd")
    assert values[1] == ("\u767d\u9053\u957f", "\u51b7\u9759", "\u4f60\u6765\u4e86")
